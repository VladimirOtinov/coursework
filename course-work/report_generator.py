import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference

class ReportGenerator:
    def __init__(self, database_manager):
        self.db_manager = database_manager

    def generate_family_income_chart(self):
        # Получаем данные для графика доходов по членам семьи
        family_income_data = self.db_manager.get_family_income_data()

        # Создаем график
        labels = [member['login'] for member in family_income_data]
        sizes = [member['income_percentage'] for member in family_income_data]

        plt.pie(sizes, labels=labels, autopct='%1.1f%%', startangle=90)
        plt.title('Доля доходов членов семьи')
        plt.savefig('family_income_chart.png')
        plt.close()

    def generate_family_category_chart(self):
        # Получаем данные для графика доходов по категориям для всей семьи
        family_category_data = self.db_manager.get_family_category_data()

        # Создаем график
        labels = [category['category'] for category in family_category_data]
        sizes = [category['category_percentage'] for category in family_category_data]

        plt.pie(sizes, labels=labels, autopct='%1.1f%%', startangle=90)
        plt.title('Доля доходов семьи по категориям')
        plt.savefig('family_category_chart.png')
        plt.close()

    def generate_individual_charts(self):
        # Получаем данные для графиков доходов и расходов по членам семьи
        members_data = self.db_manager.get_all_family_members()

        for member in members_data:
            # Генерируем график доходов
            income_data = self.db_manager.get_member_income_data(member['family_id'])
            self.generate_pie_chart(income_data, f"{member['login']}_income_chart.png", "Доходы")

            # Генерируем график расходов
            expense_data = self.db_manager.get_member_expense_data(member['family_id'])
            self.generate_pie_chart(expense_data, f"{member['login']}_expense_chart.png", "Расходы")

    def generate_pie_chart(self, data, filename, title):
        labels = [item['category'] for item in data]
        sizes = [item['percentage'] for item in data]

        fig, ax = plt.subplots()
        ax.pie(sizes, labels=labels, autopct='%1.1f%%', startangle=90)
        ax.axis('equal')  # Equal aspect ratio ensures that pie is drawn as a circle.
        plt.title(title)
        plt.savefig(filename)
        plt.close()

    def generate_excel_report(self):
        # Создаем новую книгу Excel
        wb = Workbook()

        # Генерируем диаграмму доходов
        self.generate_pie_chart_excel(wb, 'Доходы')

        # Генерируем диаграмму расходов
        self.generate_pie_chart_excel(wb, 'Расходы')

        # Сохраняем книгу
        wb.save('family_report.xlsx')

    def generate_pie_chart_excel(self, wb, chart_type):
        # Добавляем новый лист в книгу
        ws = wb.create_sheet(title=chart_type)

        # Получаем данные для диаграммы
        if chart_type == 'Доходы':
            chart_data = self.db_manager.get_family_income_data()
        else:
            chart_data = self.db_manager.get_family_expense_data()

        # Создаем диаграмму
        chart = PieChart()
        data = Reference(ws, min_col=2, min_row=1, max_col=2 + len(chart_data), max_row=1)
        labels = Reference(ws, min_col=1, min_row=2, max_row=len(chart_data) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        chart.title = f'{chart_type} семьи'

        # Добавляем диаграмму на лист
        ws.add_chart(chart, "C5")

